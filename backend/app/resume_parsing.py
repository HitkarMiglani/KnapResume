import io
import json
import re
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import docx
import pypdf

from app.config import settings

# Valid section categories
SECTIONS = ("experience", "project", "education", "skill")


class AIParsingError(RuntimeError):
    """Raised when a configured AI provider cannot parse a resume."""


def extract_text_from_pdf(content: bytes) -> str:
    """Extract text from a PDF binary content using pypdf."""
    reader = pypdf.PdfReader(io.BytesIO(content))
    pages_text = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages_text.append(text)
    return "\n".join(pages_text)


def extract_text_from_docx(content: bytes) -> str:
    """Extract text from a Word Document (.docx) binary content using python-docx.

    Extracts text from both paragraphs and table cells, joining with newlines.
    """
    doc = docx.Document(io.BytesIO(content))
    texts = []
    for para in doc.paragraphs:
        if para.text:
            texts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    texts.append(cell.text)
    return "\n".join(texts)


def extract_text(content: bytes, filename: str) -> str:
    """Extract text from binary resume based on file extension.

    Raises a ValueError if format is unsupported or filename is empty.
    """
    if not filename:
        raise ValueError("Filename cannot be empty")
    fn = filename.lower()
    if fn.endswith(".pdf"):
        return extract_text_from_pdf(content)
    elif fn.endswith(".docx"):
        return extract_text_from_docx(content)
    else:
        raise ValueError(
            f"Unsupported file format for {filename}. Only PDF and DOCX files are allowed."
        )


def clean_line_bullet(line: str) -> str:
    """Normalize resume facts by cleaning leading bullets and bullet list numbers.

    E.g., '- Bullet' -> 'Bullet', '1. Experience' -> 'Experience'.
    """
    cleaned = line.strip()
    while True:
        prev = cleaned
        # Bullet symbols (dashes, stars, bullets, circles, squares, etc.)
        cleaned = re.sub(
            r"^[\-\*\+\•\◦\▪\u2022\u2013\u2014\u25cb\u25cf\u25aa\u25ab\u25c6\ue000-\uf8ff]\s*",
            "",
            cleaned,
        )
        # Numbers/letters like "1.", "1)", "a.", "a)"
        cleaned = re.sub(r"^(?:\d+|\w)(?:\.|\))\s*", "", cleaned)
        # Parenthesized numbers like "(1)"
        cleaned = re.sub(r"^\(\d+\)\s*", "", cleaned)
        cleaned = cleaned.strip()
        if cleaned == prev:
            break
    return cleaned


def parse_resume_fallback(raw_text: str) -> list[dict[str, str]]:
    """Determine facts list deterministically using keywords and line transitions."""
    keywords_map = {
        "experience": ["experience", "work", "employment", "history", "professional"],
        "project": ["project", "portfolio"],
        "education": [
            "education",
            "academic",
            "courses",
            "school",
            "university",
            "college",
            "certification",
            "credential",
        ],
        "skill": ["skill", "technolog", "tool", "expertise", "competenc"],
    }

    current_section = "experience"
    facts = []

    for line in raw_text.splitlines():
        line_stripped = line.strip()
        if not line_stripped:
            continue

        line_lower = line_stripped.lower()

        # Check if line under 40 chars triggers transition
        is_transition = False
        if len(line_stripped) < 40:
            for sec, kws in keywords_map.items():
                if any(kw in line_lower for kw in kws):
                    current_section = sec
                    is_transition = True
                    break

        if is_transition:
            continue

        # Otherwise, clean leading bullet/numbered markers, and if non-empty, add as fact
        cleaned = clean_line_bullet(line_stripped)
        if cleaned:
            facts.append({"section": current_section, "raw_text": cleaned})

    return facts


def _parse_ai_resume_response(response: dict[str, Any]) -> list[dict[str, str]]:
    choices = response.get("choices")
    if not isinstance(choices, list) or not choices:
        raise AIParsingError("AI provider returned no choices")

    content = choices[0].get("message", {}).get("content")
    if not isinstance(content, str):
        raise AIParsingError("AI provider returned no JSON content")

    try:
        parsed = json.loads(content)
    except json.JSONDecodeError as exc:
        raise AIParsingError("AI provider returned invalid JSON") from exc

    if not isinstance(parsed, dict):
        raise AIParsingError("AI provider returned an invalid parse object")

    facts = parsed.get("facts")
    if not isinstance(facts, list):
        raise AIParsingError("AI provider did not return 'facts' array")

    cleaned_facts = []
    valid_sections = set(SECTIONS)

    for fact in facts:
        if not isinstance(fact, dict):
            raise AIParsingError("Fact item must be a dictionary")
        section = fact.get("section")
        raw_text = fact.get("raw_text")

        if not isinstance(section, str) or not isinstance(raw_text, str):
            raise AIParsingError("Fact item must contain section and raw_text as strings")

        section = section.strip().lower()
        if section not in valid_sections:
            raise AIParsingError(f"Invalid section category: {section}")

        raw_text = raw_text.strip()
        if raw_text:
            cleaned_facts.append({"section": section, "raw_text": raw_text})

    return cleaned_facts


def _parse_with_ai(raw_text: str) -> list[dict[str, str]]:
    endpoint = f"{settings.ai_base_url.rstrip('/')}/chat/completions"
    payload = {
        "model": settings.ai_model,
        "temperature": 0,
        "response_format": {"type": "json_object"},
        "messages": [
            {
                "role": "system",
                "content": (
                    "Analyze raw resume text to extract facts. "
                    "Categorize each into 'experience', 'project', 'education', or 'skill'. "
                    "Return a JSON object with a single root key 'facts' containing a "
                    "list of objects, each with keys 'section' and 'raw_text'.\n"
                    "Structure matches:\n"
                    '{"facts": [{"section": "experience"|"project"|'
                    '"education"|"skill", "raw_text": "text"}]}'
                ),
            },
            {"role": "user", "content": raw_text},
        ],
    }
    request = Request(
        endpoint,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {settings.ai_api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urlopen(request, timeout=settings.ai_timeout_seconds) as response:
            body = json.loads(response.read())
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as exc:
        raise AIParsingError("AI provider request failed") from exc

    if not isinstance(body, dict):
        raise AIParsingError("AI provider returned an invalid response")
    return _parse_ai_resume_response(body)


def parse_resume(raw_text: str) -> list[dict[str, str]]:
    """Parse resume text using LLM if configured, otherwise fallback to keyword regex."""
    if settings.ai_api_key:
        try:
            return _parse_with_ai(raw_text)
        except Exception:
            return parse_resume_fallback(raw_text)
    return parse_resume_fallback(raw_text)
